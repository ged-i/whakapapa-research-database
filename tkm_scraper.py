"""
tkm_scraper.py — populate the TKM database from tkm.govt.nz

Reads the Iwi sheet of TKM_Iwi_Hapu_Marae_Database.xlsx, fetches every iwi page,
and writes CSVs that match the workbook / SQL schema:

    out/iwi.csv, out/iwi_local_authority.csv, out/hapu.csv, out/marae.csv,
    out/marae_hapu.csv, out/representative_org.csv, out/contact.csv

Usage:
    pip install requests beautifulsoup4 openpyxl
    python tkm_scraper.py TKM_Iwi_Hapu_Marae_Database.xlsx

Be polite: the script sleeps 1 s between requests (~2 minutes for 117 pages).
Contact details are published on TKM at each organisation's discretion — respect
the TKM Disclaimer and Copyright pages, and keep this data current.
Coordinates are NOT on TKM: fill marae lat/long from the TPK ArcGIS marae layer
(GIS Marae Viewer linked at https://www.tkm.govt.nz/links/) or by geocoding.
"""
import csv, re, sys, time, pathlib
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

HEADERS = {"User-Agent": "tkm-database-builder (research use)"}
OUT = pathlib.Path("out"); OUT.mkdir(exist_ok=True)

FLAG_MAP = {  # substring on the org icon alt-text -> column
    "Treaty of Waitangi settlement negotiations": "flag_treaty_mandate",
    "Post-Treaty settlement governance entity": "flag_psge",
    "Mandated iwi organisation": "flag_mfa_mandated",
    "Recognised iwi organisation": "flag_mfa_recognised",
    "Iwi Aquaculture Organisation": "flag_iao",
    "iwi authority": "flag_rma_iwi_authority",
    "Tūhono": "flag_tuhono",
}


def text(el):
    return re.sub(r"\s+", " ", el.get_text(" ", strip=True)) if el else ""


def scrape_iwi(iwi_id, url):
    r = requests.get(url, headers=HEADERS, timeout=30)
    r.raise_for_status()
    s = BeautifulSoup(r.text, "html.parser")
    rec = {"iwi_id": iwi_id, "tkm_url": url}

    # -- population / MFA status / rohe
    body = text(s)
    m = re.search(r"Population:\s*([A-Za-z ]+\d{4})?:?\s*([\d,]+)", body)
    if m:
        rec["population_source"] = (m.group(1) or "").strip()
        rec["population"] = m.group(2).replace(",", "")
    rec["mfa_status"] = "Recognised iwi in the Māori Fisheries Act 2004" if "Recognised iwi in the Māori Fisheries Act" in body else ""
    m = re.search(r"(This rohe map represents.*?\.)", body)
    rec["rohe_description"] = m.group(1) if m else ""
    img = s.find("img", src=re.compile(r"/rohe/\d+\.jpg"))
    rec["rohe_map_image_url"] = requests.compat.urljoin(url, img["src"]) if img else ""
    doc = s.find("a", href=re.compile(r"/rohe/.*\.(pdf|jpg)", re.I))
    rec["rohe_document_url"] = requests.compat.urljoin(url, doc["href"]) if doc else ""

    # -- local authorities
    las = []
    la_head = s.find(string=re.compile("extends into the regions"))
    if la_head:
        block = la_head.find_parent().find_next("dl") or la_head.find_parent().find_next_sibling()
        current = None
        for el in block.find_all(["dt", "dd", "p", "li"]) if block else []:
            t = text(el)
            if t in ("Regional Council", "Territorial Authority"):
                current = t
            elif t and current:
                las.append((iwi_id, current, t))

    # -- hapū / marae table (blank hapū cell = same hapū as previous row)
    hapu_rows, marae_rows, links = [], [], []
    tbl = s.find("table")
    if tbl:
        prev_hapu = ""
        for tr in tbl.find_all("tr")[1:]:
            cells = [text(td) for td in tr.find_all(["td", "th"])]
            if len(cells) < 4:
                continue
            hapu, marae, wharenui, location = cells[:4]
            hapu = hapu or prev_hapu
            prev_hapu = hapu
            if hapu and hapu not in [h[1] for h in hapu_rows]:
                hapu_rows.append((f"{iwi_id}-H{len(hapu_rows)+1:03d}", hapu, iwi_id))
            hid = next(h[0] for h in hapu_rows if h[1] == hapu) if hapu else ""
            # split "60 Waiotoi Road, Ngunguru" -> address / locality
            parts = [p.strip() for p in location.rsplit(",", 1)]
            addr, loc = (parts[0], parts[1]) if len(parts) == 2 else ("", location)
            key = (marae, loc)
            existing = next((m for m in marae_rows if (m[1], m[5]) == key), None)
            if not existing:
                existing = (f"{iwi_id}-M{len(marae_rows)+1:03d}", marae, wharenui, "Iwi marae", addr, loc, url)
                marae_rows.append(existing)
            if hid:
                links.append((existing[0], hid, iwi_id))

    # -- representative organisations
    orgs, contacts = [], []
    for h3 in s.select("h3"):
        name = text(h3)
        if not name or name.lower().startswith(("this rohe", "hapū")):
            continue
        sect = []
        for sib in h3.find_next_siblings():
            if sib.name in ("h2", "h3", "hr"):
                break
            sect.append(sib)
        blob = " ".join(text(x) for x in sect)
        if "Legal entity" not in blob and "Postal Address" not in blob:
            continue
        oid = f"{iwi_id}-O{len(orgs)+1:02d}"
        o = {"org_id": oid, "iwi_id": iwi_id, "org_name": name}
        for k in FLAG_MAP.values():
            o[k] = "N"
        for img in h3.find_all_next("img", limit=12):
            alt = img.get("alt", "")
            for needle, col in FLAG_MAP.items():
                if needle in alt:
                    o[col] = "Y"
        for lbl, col in [("Legal entity", "legal_entity"), ("Governance structure", "governance_structure"),
                         ("Postal Address", "postal_address"), ("Physical Address", "physical_address"),
                         ("Phone", "phone"), ("Email", "email"), ("Website", "website"), ("Last updated", "tkm_last_updated")]:
            m = re.search(lbl + r":\s*(.*?)(?=\s(?:Legal entity|Governance structure|Chair|CEO|General Manager|Postal Address|Physical Address|Phone|Email|Website|RMA Contact|Last updated):|$)", blob)
            o[col] = m.group(1).strip() if m else ""
        m = re.search(r"\[\s*(.*?(?:Act|Settlement)[^\]]*)\]", blob)
        o["settlement_act"] = m.group(1).strip() if m else ""
        orgs.append(o)
        for role in ("Chair", "CEO", "General Manager", "Secretary", "Contact Person"):
            m = re.search(role + r":\s*([^:]+?)(?=\s(?:[A-Z][A-Za-z ]+):|$)", blob)
            if m:
                contacts.append((f"{oid}-C{len(contacts)+1:02d}", oid, role, m.group(1).strip(), "", ""))
        m = re.search(r"RMA Contact:\s*(.*?)\s*Ph:\s*([\d ]+)\s*Email:\s*(\S+)", blob)
        if m:
            contacts.append((f"{oid}-C{len(contacts)+1:02d}", oid, "RMA Contact", m.group(1), m.group(2), m.group(3)))

    return rec, las, hapu_rows, marae_rows, links, orgs, contacts


def main(xlsx):
    ws = load_workbook(xlsx, read_only=True)["Iwi"]
    iwi_list = [(r[0], r[3]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    files = {k: open(OUT / f"{k}.csv", "w", newline="", encoding="utf-8") for k in
             ("iwi", "iwi_local_authority", "hapu", "marae", "marae_hapu", "representative_org", "contact")}
    w = {k: csv.writer(f) for k, f in files.items()}
    w["iwi"].writerow(["iwi_id", "tkm_url", "mfa_status", "population", "population_source",
                       "rohe_description", "rohe_map_image_url", "rohe_document_url"])
    w["iwi_local_authority"].writerow(["iwi_id", "authority_type", "authority_name"])
    w["hapu"].writerow(["hapu_id", "hapu_name", "iwi_id"])
    w["marae"].writerow(["marae_id", "marae_name", "wharenui", "marae_type", "address_line", "locality", "tkm_url"])
    w["marae_hapu"].writerow(["marae_id", "hapu_id", "iwi_id"])
    org_cols = ["org_id", "iwi_id", "org_name", "legal_entity", "governance_structure", *FLAG_MAP.values(),
                "settlement_act", "postal_address", "physical_address", "phone", "email", "website", "tkm_last_updated"]
    w["representative_org"].writerow(org_cols)
    w["contact"].writerow(["contact_id", "org_id", "role", "full_name", "phone", "email"])

    for iwi_id, url in iwi_list:
        try:
            rec, las, hapu, marae, links, orgs, contacts = scrape_iwi(iwi_id, url)
        except Exception as e:  # keep going, report at end
            print(f"!! {iwi_id} {url}: {e}")
            continue
        w["iwi"].writerow([rec.get(c, "") for c in ("iwi_id", "tkm_url", "mfa_status", "population", "population_source",
                                                      "rohe_description", "rohe_map_image_url", "rohe_document_url")])
        w["iwi_local_authority"].writerows(las)
        w["hapu"].writerows(hapu)
        w["marae"].writerows(marae)
        w["marae_hapu"].writerows(links)
        for o in orgs:
            w["representative_org"].writerow([o.get(c, "") for c in org_cols])
        w["contact"].writerows(contacts)
        print(f"{iwi_id}: {len(hapu)} hapū, {len(marae)} marae, {len(orgs)} org(s)")
        time.sleep(1)
    for f in files.values():
        f.close()
    print("\nDone. CSVs in ./out/. NOTE: the same marae can appear under several iwi — "
          "de-duplicate marae.csv on (marae_name, locality) before loading, and re-point marae_hapu.")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "TKM_Iwi_Hapu_Marae_Database.xlsx")
