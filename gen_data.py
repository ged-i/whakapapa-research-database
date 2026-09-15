"""
gen_data.py — export the workbook's lookup tables to JSON for the web map (docs/data/*.json).

    python gen_data.py [TKM_Iwi_Hapu_Marae_Database.xlsx]

Run this after editing the spreadsheet (or after loading scraper output into it), then
commit docs/data/ — GitHub Pages redeploys automatically.
"""
import json, pathlib, sys
from openpyxl import load_workbook

xlsx = sys.argv[1] if len(sys.argv) > 1 else "TKM_Iwi_Hapu_Marae_Database.xlsx"
wb = load_workbook(xlsx, read_only=True, data_only=True)
out = pathlib.Path("docs/data"); out.mkdir(parents=True, exist_ok=True)


def rows(sheet):
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    return [dict(zip(hdr, r)) for r in it if r and r[0]]


regions = {r["region_id"]: r["region_name"] for r in rows("Region")}
iwi_regions = {}
for r in rows("Iwi_Region"):
    iwi_regions.setdefault(r["iwi_id"], []).append(regions[r["region_id"]])

orgs = {}
for o in rows("Representative_Org"):
    o = {k: ("" if v is None else v) for k, v in o.items()}
    o["contacts"] = []
    orgs[o["org_id"]] = o
for c in rows("Contact"):
    if c["org_id"] in orgs:
        orgs[c["org_id"]]["contacts"].append({k: ("" if v is None else v) for k, v in c.items()})

iwi = []
for r in rows("Iwi"):
    iwi.append({
        "id": r["iwi_id"],
        "name": r["iwi_name"],
        "regions": iwi_regions.get(r["iwi_id"], []),
        "tkm_url": r["tkm_url"] or "",
        "mfa": bool(r.get("mfa_status")),
        "population": r.get("population") or None,
        "orgs": [o for o in orgs.values() if o["iwi_id"] == r["iwi_id"]],
    })

(out / "iwi.json").write_text(json.dumps(iwi, ensure_ascii=False, indent=1), encoding="utf-8")
(out / "regions.json").write_text(json.dumps(regions, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"wrote {len(iwi)} iwi, {len(regions)} regions, {sum(len(i['orgs']) for i in iwi)} orgs -> {out}/")

# ---- whakapapa layer ----
tup = {}
if "Tupuna" in wb.sheetnames:
    for t in rows("Tupuna"):
        t = {k: ("" if v is None else v) for k, v in t.items()}
        t["links"] = []; tup[t["tupuna_id"]] = t
    if "Tupuna_Location" in wb.sheetnames:
        for l in rows("Tupuna_Location"):
            l = {k: ("" if v is None else v) for k, v in l.items()}
            if l["tupuna_id"] in tup:
                tup[l["tupuna_id"]]["links"].append(l)
(out / "tupuna.json").write_text(json.dumps(list(tup.values()), ensure_ascii=False, indent=1), encoding="utf-8")
print(f"wrote {len(tup)} tupuna")

# ---- provenance ----
src = [{k: ("" if v is None else v) for k, v in r.items()} for r in rows("Source")] if "Source" in wb.sheetnames else []
clm = [{k: ("" if v is None else v) for k, v in r.items()} for r in rows("Claim")] if "Claim" in wb.sheetnames else []
(out / "sources.json").write_text(json.dumps(src, ensure_ascii=False, indent=1), encoding="utf-8")
(out / "claims.json").write_text(json.dumps(clm, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"wrote {len(src)} sources, {len(clm)} claims")

# ---- hapū aliases ----
al = [{k: ("" if v is None else v) for k, v in r.items()} for r in rows("Hapu_Alias")] if "Hapu_Alias" in wb.sheetnames else []
(out / "hapu_aliases.json").write_text(json.dumps(al, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"wrote {len(al)} hapū aliases")
