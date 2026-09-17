"""
sync_from_json.py — rebuild the research sheets in the workbook from docs/data/*.json.

The map commits tupuna.json, claims.json and discrepancies.json to the repo directly, so for
those tables the JSON is the master record and the workbook is a report. Run this after pulling:

    python sync_from_json.py [TKM_Iwi_Hapu_Marae_Database.xlsx]
"""
import json, sys, pathlib
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
xlsx = sys.argv[1] if len(sys.argv) > 1 else "TKM_Iwi_Hapu_Marae_Database.xlsx"
wb = load_workbook(xlsx); D = pathlib.Path("docs/data")
HDR_FILL=PatternFill("solid",fgColor="1F3864"); HDR_FONT=Font(name="Arial",bold=True,color="FFFFFF",size=10); BODY=Font(name="Arial",size=10)
def write(title, cols, rows, widths=None):
    if title in wb.sheetnames: del wb[title]
    ws = wb.create_sheet(title); ws.append(cols)
    for c in range(1,len(cols)+1):
        cell=ws.cell(row=1,column=c); cell.fill=HDR_FILL; cell.font=HDR_FONT; cell.alignment=Alignment(vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width=(widths or [18]*len(cols))[c-1]
    for r in rows:
        ws.append([("" if v is None else (json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v)) for v in r])
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font=BODY; cell.alignment=Alignment(vertical="top",wrap_text=True)
    ws.freeze_panes="A2"
    if rows:
        t=Table(displayName=title,ref=f"A1:{get_column_letter(len(cols))}{len(rows)+1}"); t.tableStyleInfo=TableStyleInfo(name="TableStyleLight1",showRowStripes=True); ws.add_table(t)
    print(f"{title}: {len(rows)} rows")
tup = json.loads((D/"tupuna.json").read_text(encoding="utf-8"))
tc=["tupuna_id","name","other_names","born","died","birthplace","mother","father","notes","source"]
write("Tupuna", tc, [[t.get(c,"") for c in tc] for t in tup], [10,28,24,12,12,26,24,24,50,30])
lc=["link_id","tupuna_id","location_type","location_ref","location_detail","latitude","longitude","relationship","notes","source","as_recorded","linz_title","shares","source_id"]
write("Tupuna_Location", lc, [[l.get(c,"") if c!="tupuna_id" else t["tupuna_id"] for c in lc] for t in tup for l in t.get("links",[])], [9,10,14,30,30,11,11,16,40,28,24,12,10,9])
ic=["identity_id","tupuna_id","name","source_id","date","document_url","note"]
write("Tupuna_Identity", ic, [[x.get(c,"") if c!="tupuna_id" else t["tupuna_id"] for c in ic] for t in tup for x in t.get("identities",[])], [10,10,30,10,12,44,40])
ec=["event_id","tupuna_id","type","date","place_type","place_ref","place_detail","latitude","longitude","description","source_id","confidence","document_url","note"]
write("Tupuna_Event", ec, [[x.get(c,"") if c!="tupuna_id" else t["tupuna_id"] for c in ec] for t in tup for x in t.get("events",[])], [9,10,18,12,10,26,30,11,11,44,10,11,40,30])
dc=["disc_id","tupuna_id","title","side_a","side_b","hypothesis","status","opened","note"]
disc = json.loads((D/"discrepancies.json").read_text(encoding="utf-8")) if (D/"discrepancies.json").exists() else []
write("Discrepancy", dc, [[x.get(c,"") for c in dc] for x in disc], [9,10,40,36,36,36,10,12,30])
cc=["claim_id","source_id","subject_type","subject_ref","subject_detail","relationship","object_ref","object_detail","page_url","retrieved","confidence","note"]
clm = json.loads((D/"claims.json").read_text(encoding="utf-8"))
write("Claim", cc, [[x.get(c,"") for c in cc] for x in clm], [9,9,11,24,28,16,24,24,44,12,11,40])
wb.save(xlsx); print("saved", xlsx)
